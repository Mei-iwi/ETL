"""Request-scoped previews. Downloads continue to use the original file."""

import csv
import json
import os
import shutil
import signal
import subprocess
import tempfile
import time
from contextlib import contextmanager, nullcontext
from datetime import date, datetime
from datetime import time as day_time
from pathlib import Path
from threading import BoundedSemaphore
from xml.etree import ElementTree
from zipfile import ZipFile

import pymupdf
from openpyxl import load_workbook

from data_sync_etl.domain.core import DomainError, SourceFileError

MAX_ROWS, MAX_COLUMNS, MAX_SHEETS, MAX_CELLS = 500, 100, 20, 50000
MAX_PREVIEW_BYTES = 30 * 1024 * 1024


def check_cancelled(cancelled):
    if cancelled is not None and cancelled.is_set():
        raise SourceFileError('SOURCE_REQUEST_CANCELLED')


class LibreOfficePdfConverter:
    def __init__(self, settings):
        self.settings = settings
        self.slots = BoundedSemaphore(getattr(settings, 'source_temp_max_concurrency', 4))

    def executable(self):
        candidates = [self.settings.office_pdf_executable, shutil.which('soffice'), shutil.which('libreoffice')]
        if os.name == 'nt':
            candidates.append(Path(os.environ.get('ProgramFiles', 'C:/Program Files')) / 'LibreOffice/program/soffice.com')
        for path in candidates:
            if path and Path(path).is_file():
                return str(Path(path).resolve())
        raise SourceFileError('PREVIEW_CONVERTER_UNAVAILABLE', transient=True)

    @staticmethod
    def _stop(process):
        if process.poll() is not None:
            return
        if os.name == 'nt':
            subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           creationflags=subprocess.CREATE_NO_WINDOW, timeout=10, check=False)
        else:
            os.killpg(process.pid, signal.SIGKILL)
        process.wait(timeout=10)

    def convert(self, source, directory, cancelled=None):
        if not self.slots.acquire(blocking=False):
            raise SourceFileError('PREVIEW_BUSY', transient=True)
        try:
            return self._convert(source, directory, cancelled)
        finally:
            self.slots.release()

    def _convert(self, source, directory, cancelled=None):
        executable = self.executable()
        # Prevent the converter from resolving linked images, OLE or remote templates.
        with ZipFile(source) as archive:
            for entry in archive.infolist():
                if entry.filename.endswith('.rels'):
                    for rel in ElementTree.fromstring(archive.read(entry)):
                        if rel.get('TargetMode') == 'External' and not rel.get('Type', '').endswith('/hyperlink'):
                            raise SourceFileError('PREVIEW_EXTERNAL_CONTENT')
        work = Path(tempfile.mkdtemp(prefix='etl-preview-', dir=directory))
        profile = work / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user/registrymodifications.xcu').write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Common/Load">'
            '<prop oor:name="UpdateDocMode" oor:op="fuse"><value>3</value></prop></item></oor:items>', encoding='utf-8')
        output = work / (source.stem + '.pdf')
        args = [executable, '-env:UserInstallation=' + profile.resolve().as_uri(), '--headless', '--nologo',
                '--nodefault', '--nofirststartwizard', '--norestore', '--convert-to',
                'pdf:writer_pdf_Export' if source.suffix.lower() == '.docx' else 'pdf:impress_pdf_Export',
                '--outdir', str(work.resolve()), str(source.resolve())]
        check_cancelled(cancelled)
        process = subprocess.Popen(args, cwd=work, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                                   stderr=subprocess.DEVNULL, start_new_session=os.name != 'nt',
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        deadline = time.monotonic() + self.settings.office_pdf_timeout_seconds
        try:
            while process.poll() is None:
                check_cancelled(cancelled)
                if time.monotonic() >= deadline:
                    raise SourceFileError('PREVIEW_TIMEOUT', transient=True)
                time.sleep(0.1)
            check_cancelled(cancelled)
            if process.returncode != 0 or not output.is_file():
                raise SourceFileError('PREVIEW_CONVERSION_FAILED', transient=True)
            if output.stat().st_size > MAX_PREVIEW_BYTES:
                raise SourceFileError('FILE_TOO_LARGE')
            with pymupdf.open(output) as document:
                if not document.is_pdf or document.page_count < 1 or document.needs_pass:
                    raise SourceFileError('PREVIEW_CONVERSION_FAILED', transient=True)
            destination = directory / 'preview.pdf'
            output.replace(destination)
            return destination
        finally:
            self._stop(process)


def _cell(value, cell=None):
    text = value.isoformat() if isinstance(value, (datetime, date, day_time)) else '' if value is None else str(value)
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    result = {'text': text[:2000]}
    if cell is not None:
        font, fill = cell.font, cell.fill
        result.update(bold=bool(font and font.b), italic=bool(font and font.i))
        for key, color in [('color', font.color if font else None), ('background', fill.fgColor if fill and fill.fill_type == 'solid' else None)]:
            if color is not None and color.type == 'rgb' and isinstance(color.rgb, str):
                result[key] = '#' + color.rgb[-6:]
        if cell.alignment and cell.alignment.horizontal in ('left', 'center', 'right'):
            result['align'] = cell.alignment.horizontal
    return result


def spreadsheet_preview(path, cancelled=None):
    sheets, remaining, characters = [], MAX_CELLS, 0
    truncated = False
    if path.suffix.lower() == '.csv':
        raw = path.read_bytes()
        encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            text = raw.decode('cp1258')
        import io
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        rows = []
        for index, values in enumerate(csv.reader(io.StringIO(text, newline=''), dialect)):
            check_cancelled(cancelled)
            if index >= MAX_ROWS or characters >= 2_000_000:
                truncated = True
                break
            row = [_cell(value) for value in values[:MAX_COLUMNS]]
            characters += sum(len(cell['text']) for cell in row)
            truncated |= len(values) > MAX_COLUMNS or any(len(value) > 2000 for value in values)
            rows.append(row)
        sheets.append({'name': 'CSV', 'rows': rows})
    else:
        with path.open('rb') as stream:
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                truncated = len(workbook.worksheets) > MAX_SHEETS
                for sheet in workbook.worksheets[:MAX_SHEETS]:
                    check_cancelled(cancelled)
                    rows = []
                    columns = min(sheet.max_column or 1, MAX_COLUMNS)
                    limit = min(sheet.max_row or 1, MAX_ROWS, remaining // columns)
                    truncated |= (sheet.max_row or 1) > limit or (sheet.max_column or 1) > MAX_COLUMNS
                    if limit and characters < 2_000_000:
                        for cells in sheet.iter_rows(max_row=limit, max_col=columns):
                            check_cancelled(cancelled)
                            row = [_cell(cell.value, cell) for cell in cells]
                            rows.append(row)
                            remaining -= len(row)
                            characters += sum(len(cell['text']) for cell in row)
                            if characters >= 2_000_000:
                                truncated = True
                                break
                    sheets.append({'name': sheet.title, 'rows': rows})
            finally:
                workbook.close()
    return {'kind': 'spreadsheet', 'sheets': sheets, 'truncated': bool(truncated),
            'max_rows': MAX_ROWS, 'max_columns': MAX_COLUMNS}


class FilePreviewService:
    def __init__(self, resolver, settings):
        self.resolver = resolver
        self.converter = LibreOfficePdfConverter(settings)

    @contextmanager
    def temporary_path(self, reference, *, cancelled=None):
        with self.resolver.temporary_path(reference, cancelled=cancelled) as source:
            suffix = source.suffix.lower()
            try:
                validator = self.resolver.formats.find_validator(reference.filename)
                size = source.stat().st_size
                with source.open('rb') as stream:
                    validator.validate(suffix, reference.mime, stream.read(16), size)
                validator.validate_path(source)
            except DomainError:
                raise SourceFileError('PREVIEW_INVALID_FILE') from None
            if suffix not in ('.docx', '.pptx', '.xlsx', '.csv'):
                yield source
                return
            # SOURCE derivatives share the source lease, limits and cleanup.
            scope = nullcontext(source.parent) if reference.ownership == 'SOURCE_BACKED' else tempfile.TemporaryDirectory(prefix='etl-local-preview-')
            with scope as directory:
                directory = Path(directory)
                try:
                    check_cancelled(cancelled)
                    revision = self.resolver._revision(reference, cancelled) if reference.ownership == 'SOURCE_BACKED' else None
                    if suffix in ('.docx', '.pptx'):
                        preview = self.converter.convert(source, directory, cancelled)
                    else:
                        preview = directory / 'preview.json'
                        preview.write_text(json.dumps(spreadsheet_preview(source, cancelled), ensure_ascii=False), encoding='utf-8')
                    check_cancelled(cancelled)
                    if reference.ownership == 'SOURCE_BACKED':
                        self.resolver._recheck(reference)
                        if revision != self.resolver._revision(reference, cancelled):
                            raise SourceFileError('SOURCE_REVISION_CHANGED')
                    yield preview
                except SourceFileError:
                    raise
                except (DomainError, ValueError, OSError, csv.Error, ElementTree.ParseError):
                    raise SourceFileError('PREVIEW_INVALID_FILE') from None
