"""Bounded native document extraction; one existing task per non-PDF document.

Task page_num is an internal work index, not a physical page. Character spans
in result_metadata preserve provenance without duplicating text or adding tables.
"""

import csv
import zipfile
from datetime import date, datetime
from io import BytesIO, TextIOWrapper
from time import monotonic

from data_sync_etl.domain.core import DomainError
from data_sync_etl.ports.processing import OCRResult

MAX_UNITS = 5000
MAX_TEXT_CHARS = 1_000_000
MAX_UNIT_CHARS = 16000
MAX_ROWS = 100000
MAX_CELLS = 100000
MAX_COLUMNS = 256
MAX_SHEETS = 50
MAX_SECONDS = 60


class NativeDocumentExtractor:
    @staticmethod
    def recognize_image(path, output, engine):
        return recognize_image(path, output, engine)

    def extract(self, path):
        extension = path.suffix.lower()
        parsers = {'.docx': self.docx, '.xlsx': self.xlsx, '.txt': self.txt,
                   '.csv': self.csv, '.pptx': self.pptx}
        if extension not in parsers:
            raise DomainError('UNSUPPORTED_PROCESSING_FORMAT')
        texts, units, offset = [], [], 0
        started = monotonic()
        try:
            for text, title, locator in parsers[extension](path):
                if monotonic() - started > MAX_SECONDS:
                    raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                if not text.strip():
                    continue
                for start in range(0, len(text), MAX_UNIT_CHARS):
                    part = text[start:start + MAX_UNIT_CHARS]
                    if not part.strip():
                        continue
                    if len(units) >= MAX_UNITS or offset + len(part) > MAX_TEXT_CHARS:
                        raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                    units.append({'start': offset, 'end': offset + len(part),
                                  'title': title[:500], 'locator': {**locator, 'char_offset': start}})
                    texts.append(part)
                    offset += len(part) + 2
            if not texts:
                raise DomainError('EMPTY_EXTRACTED_CONTENT')
            return OCRResult('\n\n'.join(texts), engine_name='native_' + extension[1:],
                             metadata={'source_type': extension[1:], 'logical_units': units,
                                       **({'structure_version': 'docx-outline-v1'} if extension == '.docx' else {})})
        except DomainError:
            raise
        except UnicodeError:
            raise DomainError('UNSUPPORTED_TEXT_ENCODING') from None
        except Exception:
            raise DomainError('INVALID_DOCUMENT_CONTENT') from None

    def extract_bytes(self, filename, data, *, max_uncompressed_bytes=128 * 1024 * 1024):
        """Extract supported source documents from bounded RAM without staging a file."""
        from pathlib import PurePosixPath
        extension = PurePosixPath(filename).suffix.lower()
        parsers = {'.docx': self.docx, '.xlsx': self.xlsx, '.txt': self.txt,
                   '.csv': self.csv, '.pptx': self.pptx}
        if extension not in parsers:
            raise DomainError('UNSUPPORTED_PROCESSING_FORMAT')
        if extension in ('.docx', '.xlsx', '.pptx'):
            try:
                with zipfile.ZipFile(BytesIO(data)) as archive:
                    entries = archive.infolist()
                    expanded = sum(item.file_size for item in entries)
                    if (len(entries) > 10000 or expanded > max_uncompressed_bytes
                            or any(item.file_size > 32 * 1024 * 1024
                                   or item.file_size > 1024 * max(item.compress_size, 1)
                                   for item in entries)):
                        raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
            except DomainError:
                raise
            except (OSError, zipfile.BadZipFile):
                raise DomainError('INVALID_OPENXML_CONTAINER') from None
        texts, units, offset = [], [], 0
        started = monotonic()
        try:
            for text, title, locator in parsers[extension](BytesIO(data)):
                if monotonic() - started > MAX_SECONDS:
                    raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                if not text.strip():
                    continue
                for start in range(0, len(text), MAX_UNIT_CHARS):
                    part = text[start:start + MAX_UNIT_CHARS]
                    if not part.strip():
                        continue
                    if len(units) >= MAX_UNITS or offset + len(part) > MAX_TEXT_CHARS:
                        raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                    units.append({'start': offset, 'end': offset + len(part),
                                  'title': title[:500], 'locator': {**locator, 'char_offset': start}})
                    texts.append(part)
                    offset += len(part) + 2
            if not texts:
                raise DomainError('EMPTY_EXTRACTED_CONTENT')
            return OCRResult('\n\n'.join(texts), engine_name='native_' + extension[1:],
                             metadata={'source_type': extension[1:], 'logical_units': units,
                                       **({'structure_version': 'docx-outline-v1'} if extension == '.docx' else {})})
        except DomainError:
            raise
        except UnicodeError:
            raise DomainError('UNSUPPORTED_TEXT_ENCODING') from None
        except Exception:
            raise DomainError('INVALID_DOCUMENT_CONTENT') from None

    @staticmethod
    def docx(path):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        for index, block in enumerate(document.iter_inner_content(), 1):
            if index > MAX_UNITS:
                raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
            if isinstance(block, Paragraph):
                # Trust OOXML outline levels (including inherited styles), not text regex.
                level, style, visited = None, block.style, set()
                paragraph_properties = block._p.pPr
                properties = [paragraph_properties] if paragraph_properties is not None else []
                while style is not None and style.style_id not in visited:
                    visited.add(style.style_id)
                    if style.element.pPr is not None:
                        properties.append(style.element.pPr)
                    style = style.base_style
                for props in properties:
                    outline = props.find(qn('w:outlineLvl'))
                    if outline is not None:
                        value = outline.get(qn('w:val'))
                        level = int(value) + 1 if value is not None and value.isdigit() else 0
                        break
                yield block.text, f'Word · Đoạn {index}', {'block': index, 'kind': 'paragraph',
                                                        'style': block.style.name if block.style else '',
                                                        **({'heading_level': level} if level is not None and level != 10 else {})}
            elif isinstance(block, Table):
                for row_index, row in enumerate(block.rows, 1):
                    if row_index > MAX_ROWS or len(row.cells) > MAX_COLUMNS:
                        raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                    cells = [cell.text for cell in row.cells]
                    if not any(cell.strip() for cell in cells):
                        continue
                    text = ' | '.join(cells)
                    yield text, f'Word · Bảng {index}, hàng {row_index}', {
                        'block': index, 'kind': 'table', 'row': row_index}

    @staticmethod
    def xlsx(path):
        from openpyxl import load_workbook
        if hasattr(path, 'getvalue'):
            raw = path.getvalue()
            cached_input, formula_input = BytesIO(raw), BytesIO(raw)
        else:
            cached_input = formula_input = path
        cached = load_workbook(cached_input, read_only=True, data_only=True, keep_links=False)
        formulas = None
        try:
            formulas = load_workbook(formula_input, read_only=True, data_only=False, keep_links=False)
            if len(cached.worksheets) > MAX_SHEETS:
                raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
            cell_count = 0
            for sheet, formula_sheet in zip(cached.worksheets, formulas.worksheets, strict=True):
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                for row_no, (row, formula_row) in enumerate(zip(sheet.rows, formula_sheet.rows, strict=True), 1):
                    cell_count += len(row)
                    if row_no > MAX_ROWS or len(row) > MAX_COLUMNS or cell_count > MAX_CELLS:
                        raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                    values, coordinates, has_value = [], [], False
                    for cell, formula in zip(row, formula_row, strict=True):
                        value = cell.value
                        if value is not None and str(value).strip():
                            has_value = True
                        if value is None and formula.data_type == 'f':
                            value = '[Công thức chưa có giá trị cache]'
                        if value is None:
                            continue
                        if isinstance(value, (datetime, date)):
                            value = value.isoformat()
                        value = str(value)
                        if len(value) > MAX_UNIT_CHARS:
                            raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                        coordinates.append(formula.coordinate)
                        values.append(f'{formula.coordinate}: {value}')
                    if values and has_value:
                        yield ' | '.join(values), f'Excel · {sheet.title}, hàng {row_no}', {
                            'sheet': sheet.title, 'row': row_no, 'cells': coordinates}
        finally:
            cached.close()
            if formulas is not None:
                formulas.close()
            if hasattr(path, 'getvalue'):
                cached_input.close()
                formula_input.close()

    @staticmethod
    def txt(path):
        stream = TextIOWrapper(path if hasattr(path, 'read') else path.open('rb'), encoding='utf-8-sig', newline='')
        try:
            for line_no, line in enumerate(stream, 1):
                if line_no > MAX_ROWS:
                    raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                yield line, f'TXT · Dòng {line_no}', {'line': line_no}
        finally:
            stream.detach()

    @staticmethod
    def csv(path):
        stream = TextIOWrapper(path if hasattr(path, 'read') else path.open('rb'), encoding='utf-8-sig', newline='')
        try:
            sample = stream.read(8192)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(stream, dialect, strict=True)
            previous_line = 0
            try:
                for row_no, row in enumerate(reader, 1):
                    if row_no > MAX_ROWS or len(row) > MAX_COLUMNS or any(len(cell) > MAX_UNIT_CHARS for cell in row):
                        raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                    line_from = previous_line + 1
                    previous_line = reader.line_num
                    if not any(cell.strip() for cell in row):
                        continue
                    yield ' | '.join(f'C{index}: {value}' for index, value in enumerate(row, 1)), \
                        f'CSV · Hàng {row_no}, dòng {line_from}–{reader.line_num}', {
                            'row': row_no, 'line_from': line_from, 'line_to': reader.line_num,
                            'delimiter': dialect.delimiter}
            except csv.Error:
                raise DomainError('INVALID_CSV') from None
        finally:
            stream.detach()

    @staticmethod
    def pptx(path):
        from pptx import Presentation
        document = Presentation(path)
        for slide_no, slide in enumerate(document.slides, 1):
            if slide_no > 500:
                raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
            # Stable top-to-bottom/left-to-right order, retaining shape identity.
            for shape in sorted(slide.shapes, key=lambda shape: (shape.top, shape.left, shape.shape_id)):
                if shape.has_text_frame:
                    yield shape.text, f'PowerPoint · Slide {slide_no}, khối {shape.shape_id}', {
                        'slide': slide_no, 'shape_id': shape.shape_id}
                elif shape.has_table:
                    for row_no, row in enumerate(shape.table.rows, 1):
                        if row_no > MAX_ROWS or len(row.cells) > MAX_COLUMNS:
                            raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
                        cells = [cell.text for cell in row.cells]
                        if not any(cell.strip() for cell in cells):
                            continue
                        yield ' | '.join(cells), \
                            f'PowerPoint · Slide {slide_no}, bảng {shape.shape_id}, hàng {row_no}', {
                                'slide': slide_no, 'shape_id': shape.shape_id, 'row': row_no}


def recognize_image(path, output, engine):
    from PIL import Image, ImageOps
    if not getattr(engine, 'supports_images', False):
        raise DomainError('BLOCKED_OCR_RUNTIME')
    try:
        with Image.open(path) as image:
            with ImageOps.exif_transpose(image) as oriented:
                with oriented.convert('RGB') as rgb:
                    rgb.save(output, format='PNG')
        result = engine.recognize(output)
        if not result.text.strip():
            raise DomainError('EMPTY_EXTRACTED_CONTENT')
        if len(result.text) > MAX_TEXT_CHARS:
            raise DomainError('EXTRACTION_LIMIT_EXCEEDED')
        return OCRResult(result.text, confidence=result.confidence, engine_name=result.engine_name,
                         engine_version=result.engine_version, metadata={
                             **result.metadata, 'source_type': 'image',
                             'logical_units': [{'start': start, 'end': min(start + MAX_UNIT_CHARS, len(result.text)),
                                                'title': 'Ảnh · OCR chữ',
                                                'locator': {'kind': 'image', 'char_offset': start}}
                                               for start in range(0, len(result.text), MAX_UNIT_CHARS)],
                         }), str(output)
    except Exception:
        output.unlink(missing_ok=True)
        raise
