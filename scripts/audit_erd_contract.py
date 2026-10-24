"""Static ERD/Excel/SQL/ORM audit. Never loads .env, connects to DB, or runs migrations."""
import argparse
import ast
import csv
import hashlib
import io
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'TAI_LIEU_DOI_CHIEU'
OUTPUT = ROOT / '.test-tmp/audits/erd-contract'
FILES = (
    'ERD_SOURCE_GOC_15_BANG.jpg', 'ERD_MASTER_DE_XUAT_10_BANG.png',
    'MO_TA_SOURCE_POSTGRESQL.xlsx', 'MO_TA_MONGODB_BOW.xlsx',
    'QUY_TAC_TAO_MA_POSTGRESQL.xlsx', 'BAO_CAO_KIEM_TOAN_2026_10_17.md',
    'AUDIT_SOURCE_READONLY.sql', 'AUDIT_MASTER_READONLY.sql',
)
BUSINESS = {
    'master_learning_resources', 'master_users', 'master_grade_levels', 'subjects',
    'resource_subjects', 'resource_grade_levels', 'resource_versions', 'content_units',
    'ocr_jobs', 'ocr_page_tasks',
}


def read(path):
    return (ROOT / path).read_text(encoding='utf-8-sig')


def line(path, token):
    return next((n for n, value in enumerate(read(path).splitlines(), 1) if token in value), 1)


def evidence(path, token):
    return f'{path}:{line(path, token)}'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def readonly_guard(event, args):
    """Fail closed on accidental dotenv reads or network/database connections."""
    if event == 'open' and args and isinstance(args[0], str) and Path(args[0]).name.lower() == '.env':
        raise RuntimeError('Audit must not open .env')
    if event == 'socket.connect':
        raise RuntimeError('Audit must not connect to network/database services')


def split_sql(value):
    parts, start, depth, quote = [], 0, 0, False
    for pos, char in enumerate(value):
        if char == "'":
            quote = not quote
        elif not quote:
            depth += (char == '(') - (char == ')')
            if char == ',' and depth == 0:
                parts.append(value[start:pos].strip())
                start = pos + 1
    parts.append(value[start:].strip())
    return parts


def source_schema():
    path = 'alembic_source/schema_v1.sql'
    sql = read(path)
    tables = {}
    for match in re.finditer(r'CREATE TABLE (\w+)\s*\((.*?)\);', sql, re.S):
        table, body = match.groups()
        entries = split_sql(body)
        composite = next((re.findall(r'\b\w+\b', x[x.index('(') + 1:x.rindex(')')])
                          for x in entries if x.startswith('PRIMARY KEY')), [])
        cols = {}
        for entry in entries:
            found = re.match(r'(\w+)\s+(VARCHAR\(\d+\)|\w+)', entry)
            if not found or found[1] in {'PRIMARY', 'CHECK', 'UNIQUE', 'CONSTRAINT', 'FOREIGN'}:
                continue
            name, kind = found.groups()
            fk = re.search(r'REFERENCES (\w+)\((\w+)\)', entry)
            inline_pk = 'PRIMARY KEY' in entry
            cols[name] = {
                'type': kind.upper(), 'nullable': not ('NOT NULL' in entry or inline_pk or name in composite),
                'pk': inline_pk or name in composite,
                'unique': 'UNIQUE' in entry or inline_pk,
                'fk': f'{fk[1]}.{fk[2]}' if fk else '',
                'definition': entry,
                'evidence': f'{path}:{sql[:match.start()].count(chr(10)) + 1}',
            }
        tables[table] = {'columns': cols, 'constraints': [x for x in entries if x.startswith(('PRIMARY', 'CHECK', 'UNIQUE', 'CONSTRAINT', 'FOREIGN'))]}
    enums = {m[1]: re.findall(r"'([^']+)'", m[2])
             for m in re.finditer(r'CREATE TYPE (\w+) AS ENUM \((.*?)\);', sql)}
    return tables, enums


def workbook(path):
    from openpyxl import load_workbook
    return load_workbook(REFERENCE / path, read_only=True, data_only=False)


def field_rows(sheet):
    for row in sheet.iter_rows(min_row=4):
        cells = [cell.value for cell in row[:9]]
        if (len(cells) >= 6 and isinstance(cells[0], str)
                and re.fullmatch(r'[a-z_][a-z_0-9]*', cells[0]) and cells[1] is not None):
            yield row[0].row, cells


def flag(value):
    return str(value).upper() == 'TRUE'


def expected_type(kind, size):
    if str(kind).upper() in ('VARCHAR', 'CHAR'):
        width = re.search(r'\d+', str(size))
        return str(kind).upper() + (f'({width[0]})' if width else '')
    return str(kind).upper()


def source_matrix(tables, enums):
    book = workbook('MO_TA_SOURCE_POSTGRESQL.xlsx')
    rows, enum_rows, relations = [], [], []
    for table, data in tables.items():
        found = set()
        for rownum, values in field_rows(book[table]):
            name, kind, size, nullable, pk, constraints, convention, _description, _notes = values
            found.add(name)
            current = data['columns'].get(name)
            fks = re.findall(r'FK\s*→\s*(\w+\.\w+)', str(constraints or ''))
            expected = {'type': expected_type(kind, size), 'nullable': flag(nullable), 'pk': flag(pk), 'fk': ','.join(fks)}
            differences = ['MISSING_SQL_FIELD'] if current is None else [
                key for key, val in expected.items() if val != current[key]
            ]
            if current and 'UNIQUE' in str(constraints) and not current['unique']:
                differences.append('unique')
            rows.append({'table': table, 'field': name, 'ERD_Excel': json.dumps(expected, ensure_ascii=False),
                         'SQL_current': json.dumps(current, ensure_ascii=False) if current else '',
                         'difference_proven': ','.join(differences) or 'MATCH_STATIC',
                         'excel_evidence': f'TAI_LIEU_DOI_CHIEU/MO_TA_SOURCE_POSTGRESQL.xlsx#{table}!A{rownum}:I{rownum}',
                         'erd_evidence': f'ERD_SOURCE_GOC_15_BANG.jpg/table={table}',
                         'convention': convention or '', 'constraints_excel': constraints or '',
                         'scope': 'SOURCE_PRESERVE_NO_DDL'})
        for name, current in data['columns'].items():
            if name not in found:
                rows.append({'table': table, 'field': name, 'ERD_Excel': '', 'SQL_current': json.dumps(current),
                             'difference_proven': 'SQL_FIELD_NOT_IN_EXCEL', 'excel_evidence': table,
                             'erd_evidence': f'ERD_SOURCE_GOC_15_BANG.jpg/table={table}', 'convention': '',
                             'constraints_excel': '', 'scope': 'REVIEW_NO_DDL'})
        for name, current in data['columns'].items():
            if current['fk']:
                relations.append({'table': table, 'column': name, 'target': current['fk'],
                                  'direction': 'child -> parent', 'cardinality': 'N:1; UNIQUE may narrow to 1:1',
                                  'nullable': current['nullable'], 'unique_child': current['unique'],
                                  'on_delete': 'RESTRICT', 'evidence': current['evidence']})
    expected_enums = {}
    for row in book['Enums'].iter_rows(min_row=5, values_only=True):
        if row[0] in enums and row[1] is not None:
            expected_enums.setdefault(row[0], []).append(row[1])
    for name, actual in enums.items():
        enum_rows.append({'enum': name, 'excel': '|'.join(expected_enums.get(name, [])),
                          'sql': '|'.join(actual),
                          'status': 'MATCH_STATIC' if expected_enums.get(name) == actual else 'REVIEW_CONFLICT'})
    book.close()
    return rows, enum_rows, relations


def orm_schema():
    sys.path.insert(0, str(ROOT / 'src'))
    from sqlalchemy.dialects import postgresql

    from data_sync_etl.db.models import Base
    path = 'src/data_sync_etl/db/models.py'
    tree = ast.parse(read(path))
    locations = {}
    for node in tree.body:
        if not isinstance(node, ast.ClassDef):
            continue
        table = next((x.value.value for x in node.body if isinstance(x, ast.Assign)
                      and any(isinstance(t, ast.Name) and t.id == '__tablename__' for t in x.targets)), None)
        if table:
            locations[table] = {x.targets[0].id: x.lineno for x in node.body if isinstance(x, ast.Assign)
                                and isinstance(x.targets[0], ast.Name)}
    result = {}
    for name, table in Base.metadata.tables.items():
        result[name] = {
            'columns': {c.name: {'type': c.type.compile(dialect=postgresql.dialect()), 'nullable': c.nullable,
                                  'pk': c.primary_key, 'fk': sorted(f.target_fullname for f in c.foreign_keys),
                                  'evidence': f'{path}:{locations[name].get(c.name, 1)}'} for c in table.columns},
            'constraints': sorted(({'kind': type(c).__name__, 'name': c.name,
                                    'columns': [x.name for x in c.columns],
                                    'expression': str(c.sqltext) if hasattr(c, 'sqltext') else ''}
                                   for c in table.constraints), key=lambda x: (
                                       x['kind'], x['name'] or '', ','.join(x['columns']), x['expression'])),
            'indexes': [{'name': i.name, 'columns': [c.name for c in i.columns], 'unique': i.unique}
                        for i in sorted(table.indexes, key=lambda x: x.name)],
        }
    return result


def source_for_master(table, column):
    if column in ('source_origin', 'source_state_hash', 'source_missing_at'):
        return '', 'ETL governance; source authority, not user-editable source metadata'
    if table not in BUSINESS:
        return '', 'ETL runtime; not a SOURCE table copy'
    if table == 'master_learning_resources':
        if column == 'provider_type':
            return 'providers.provider_type', 'SOURCE provider_type -> DTO type -> Master provider_type'
        if column.startswith('provider_'):
            return 'providers.' + ('id' if column == 'provider_id' else column[9:]), 'denormalized provider snapshot'
        if column.startswith('resource_type_'):
            return 'resource_types.' + ('id' if column == 'resource_type_id' else column[14:]), 'denormalized resource type snapshot'
        return 'learning_resources.' + column, 'SOURCE metadata; preserve ownership'
    if table == 'master_users':
        if column.startswith('role_'):
            return 'roles.' + ('id' if column == 'role_id' else column[5:]), 'denormalized role snapshot'
        if column in ('deleted_at', 'deleted_by'):
            return '', 'Master extension; absent from SOURCE users'
        return 'users.' + column, 'SOURCE identity; hashed_password is sensitive, never Mongo metadata'
    if table == 'master_grade_levels':
        if column == 'id':
            return 'grade_levels.id', 'grade identity'
        if column == 'education_level_id':
            return 'grade_levels.education_level_id', 'logical catalogue reference after merge'
        if column.startswith('grade_'):
            return 'grade_levels.' + column[6:], 'denormalized grade snapshot'
        return 'education_levels.' + column[10:], 'denormalized education snapshot'
    if table in ('subjects', 'resource_subjects', 'resource_grade_levels'):
        return table + '.' + column, 'SOURCE metadata / N:N links'
    if table == 'resource_versions':
        if column in ('resource_id', 'version_number', 'source_hash'):
            return 'resource_versions.' + column, 'source version -> imported target version lineage'
        return '', 'ETL target version; generated ID/content-addressed storage/lifecycle; workflow fields not copied by import'
    return '', 'ETL-generated OCR/content metadata; SOURCE OCR history is not imported'


def master_matrix(orm, tables):
    result = []
    for table, data in orm.items():
        for name, col in data['columns'].items():
            source, meaning = source_for_master(table, name)
            source_col = None
            if source:
                st, sc = source.split('.')
                source_col = tables.get(st, {}).get('columns', {}).get(sc)
            result.append({'master_table': table, 'master_field': name, 'group': 'ERD_10' if table in BUSINESS else 'RUNTIME_8',
                           'source_field': source, 'source_contract': json.dumps(source_col, ensure_ascii=False) if source_col else '',
                           'master_contract': json.dumps(col, ensure_ascii=False), 'meaning': meaning,
                           'action': 'KEEP_RUNTIME' if table not in BUSINESS else 'KEEP_NO_DDL_REVIEW_MASTER_IMAGE',
                           'erd_evidence': f'ERD_MASTER_DE_XUAT_10_BANG.png/table={table}' if table in BUSINESS else 'Not drawn; runtime extension',
                           'mapping_evidence': evidence('src/data_sync_etl/adapters/source/aggregate.py', 'def build_aggregate')})
    return result


def mongo_matrix():
    from data_sync_etl.adapters.mongo.schema import INDEXES, SCHEMA_VERSION, contracts
    contract = contracts(384)
    book = workbook('MO_TA_MONGODB_BOW.xlsx')
    rows = []
    for name, validator in contract.items():
        props = validator['$jsonSchema']['properties']
        names = set()
        for rownum, values in field_rows(book[name]):
            field, kind, size, nullable, _pk, _constraints, convention, _desc, notes = values
            if rownum >= 13 and name == 'learning_resources':
                # This is an embedded author schema, not extra top-level collection fields.
                field = 'authors[].' + field
                actual = props['authors']['items']['properties'].get(field.split('.')[-1])
            else:
                names.add(field)
                actual = props.get(field)
            rows.append({'collection': name, 'field': field, 'excel_type': kind, 'excel_size': size,
                         'excel_nullable': flag(nullable), 'excel_convention': convention or '',
                         'validator': json.dumps(actual, ensure_ascii=False),
                         'status': 'MISSING_VALIDATOR_FIELD' if actual is None else
                                   'REVIEW_FALLBACK_ENUM' if name == 'content_bags' and field == 'unit_type' else 'PRESENT_REVIEW_CONSTRAINTS',
                         'excel_evidence': f'TAI_LIEU_DOI_CHIEU/MO_TA_MONGODB_BOW.xlsx#{name}!A{rownum}:I{rownum}',
                         'code_evidence': evidence('src/data_sync_etl/adapters/mongo/schema.py', f'"{name}":'),
                         'notes': notes or ''})
        for field in sorted(set(props) - names):
            rows.append({'collection': name, 'field': field, 'excel_type': '', 'excel_size': '',
                         'excel_nullable': '', 'excel_convention': '', 'validator': json.dumps(props[field]),
                         'status': 'RUNTIME_EXTENSION_NOT_IN_EXCEL', 'excel_evidence': name,
                         'code_evidence': evidence('src/data_sync_etl/adapters/mongo/schema.py', f'"{field}":'),
                         'notes': 'Existing runtime/fencing field. Preserve, document, do not remove.'})
    book.close()
    indexes = {name: [{'name': i.name, 'keys': i.keys, 'unique': i.unique} for i in values]
               for name, values in INDEXES.items()}
    return rows, {'schema_version': SCHEMA_VERSION, 'dimension': 384, 'validators': contract, 'indexes': indexes}


def migration_inventory():
    rows, chains = [], {}
    for label, folder in [('SOURCE', 'alembic_source/versions'), ('MASTER', 'alembic/versions')]:
        revisions = {}
        for path in sorted((ROOT / folder).glob('*.py')):
            if path.name == '__init__.py':
                continue
            text = path.read_text(encoding='utf-8')
            tree = ast.parse(text)
            values = {}
            for node in tree.body:
                if isinstance(node, ast.Assign) and isinstance(node.targets[0], ast.Name):
                    try:
                        values[node.targets[0].id] = ast.literal_eval(node.value)
                    except (ValueError, TypeError):
                        pass
            revision, parent = values.get('revision'), values.get('down_revision')
            if revision:
                revisions[revision] = parent
                rows.append({'layer': label, 'kind': 'REVISION', 'path': path.relative_to(ROOT).as_posix(),
                             'revision': revision, 'down_revision': parent or '', 'upgrade': any(isinstance(x, ast.FunctionDef) and x.name == 'upgrade' for x in tree.body),
                             'downgrade': any(isinstance(x, ast.FunctionDef) and x.name == 'downgrade' for x in tree.body),
                             'preflight_mention': 'preflight' in text.lower(), 'status': 'STATIC_ONLY_NOT_APPLIED',
                             'ddl_execution': 'NOT_RUN', 'evidence': f'{path.relative_to(ROOT).as_posix()}:1'})
        parents = set(revisions.values()) - {None}
        heads = sorted(set(revisions) - parents)
        chains[label] = {'revisions': revisions, 'heads': heads,
                         'single_chain': len(heads) == 1 and list(revisions.values()).count(None) == 1
                         and all(p is None or p in revisions for p in revisions.values())
                         and all(n == 1 for p, n in Counter(revisions.values()).items())}
    for path in ['alembic_source/schema_v1.sql', 'alembic_source/triggers_v1.sql',
                 'alembic_source/timestamps_v2.sql', 'scripts/sql/source_reader.sql',
                 'scripts/sql/verify_source.sql', 'scripts/sql/verify_master.sql',
                 'TAI_LIEU_DOI_CHIEU/AUDIT_SOURCE_READONLY.sql', 'TAI_LIEU_DOI_CHIEU/AUDIT_MASTER_READONLY.sql']:
        sql = read(path)
        cleaned = re.sub(r'--[^\n]*', '', sql)
        statements = [x.strip() for x in cleaned.split(';') if x.strip()]
        readonly = bool(statements) and all(re.match(r'SELECT\b', x, re.I) for x in statements)
        rows.append({'layer': 'SOURCE' if 'source' in path.lower() else 'MASTER', 'kind': 'SQL',
                     'path': path, 'revision': '', 'down_revision': '', 'upgrade': '', 'downgrade': '',
                     'preflight_mention': False, 'status': 'SELECT_ONLY_REVIEWED' if readonly else 'WRITES_PRESENT_DO_NOT_EXECUTE',
                     'ddl_execution': 'NOT_RUN', 'evidence': f'{path}:1'})
    return rows, chains


def csv_text(rows):
    stream = io.StringIO(newline='')
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]), lineterminator='\n')
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue()


def trigger_inventory():
    source_path = 'alembic_source/versions/s0002_triggers.py'
    tree = ast.parse(read(source_path))
    timestamp_tables = next(ast.literal_eval(n.value) for n in tree.body
                            if isinstance(n, ast.Assign) and isinstance(n.targets[0], ast.Name)
                            and n.targets[0].id == 'TIMESTAMP_TABLES')
    rows = []

    def add(layer, table, suffix, function, event, path, revision):
        rows.append({'layer': layer, 'table': table, 'trigger': 'trg_' + table + suffix,
                     'function': function, 'event': event, 'installed_in': revision,
                     'removed_in': revision + ' downgrade', 'evidence': evidence(path, 'def upgrade'),
                     'live': 'NOT_VERIFIED'})

    for table in timestamp_tables:
        add('SOURCE', table, '_updated_at', 'source_set_updated_at', 'BEFORE UPDATE', source_path, 's0002')
    for table in timestamp_tables[:8]:
        add('SOURCE', table, '_soft_delete_only', 'source_require_soft_delete', 'BEFORE DELETE', source_path, 's0002')
    for table in ('resource_subjects', 'resource_grade_levels'):
        add('SOURCE', table, '_touch_resource', 'source_touch_resource', 'AFTER INSERT/UPDATE/DELETE', source_path, 's0002')
    for table in ('providers', 'resource_types', 'roles', 'education_levels'):
        add('SOURCE', table, '_touch_dependents', 'source_touch_dependents', 'AFTER UPDATE', source_path, 's0002')
    master_path = 'alembic/versions/0006_source_compatibility.py'
    for table in ('resource_versions', 'ocr_jobs', 'ocr_page_tasks', 'content_units', 'ocr_page_results'):
        add('MASTER', table, '_local_updated_at', 'master_set_local_updated_at', 'BEFORE UPDATE', master_path, '0006')
    add('MASTER', 'content_units', '', 'master_guard_content_parent',
        'BEFORE INSERT/UPDATE OF parent_unit_id,resource_version_id', master_path, '0006')
    rows[-1]['trigger'] = 'trg_content_parent_guard'
    return rows


def endpoint_inventory():
    path = 'src/data_sync_etl/api/routes.py'
    rows = []
    for node in ast.parse(read(path)).body:
        if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        for dec in node.decorator_list:
            if (isinstance(dec, ast.Call) and isinstance(dec.func, ast.Attribute)
                    and isinstance(dec.func.value, ast.Name) and dec.func.value.id == 'router'):
                rows.append({'method': dec.func.attr.upper(), 'path': ast.literal_eval(dec.args[0]),
                             'handler': node.name, 'evidence': f'{path}:{node.lineno}',
                             'live': 'NOT_RUN', 'flow': 'See endpoint_flow_matrix.md'})
    return rows


def collect():
    for name in FILES:
        if not (REFERENCE / name).is_file():
            raise SystemExit('Required reference missing: ' + name)
    tables, enums = source_schema()
    source, enum_rows, relations = source_matrix(tables, enums)
    orm = orm_schema()
    master = master_matrix(orm, tables)
    mongo, mongo_snapshot = mongo_matrix()
    migrations, chains = migration_inventory()
    triggers = trigger_inventory()
    endpoints = endpoint_inventory()
    snapshot = {
        'scope': 'STATIC_ONLY; NO_DOTENV; NO_DATABASE; NO_DDL; NO_MIGRATION_EXECUTION',
        'baseline_head': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        'references': {name: sha(REFERENCE / name) for name in FILES},
        'implementation_sha256': {
            path.relative_to(ROOT).as_posix(): sha(path)
            for path in sorted({
                ROOT / 'alembic_source/schema_v1.sql', ROOT / 'alembic_source/triggers_v1.sql',
                ROOT / 'alembic_source/timestamps_v2.sql', ROOT / 'alembic/schema_v1.py',
                ROOT / 'src/data_sync_etl/db/models.py',
                ROOT / 'src/data_sync_etl/adapters/mongo/schema.py',
                *ROOT.glob('alembic/versions/*.py'), *ROOT.glob('alembic_source/versions/*.py'),
            })
        },
        'source': tables, 'source_enums': enums, 'master': orm, 'mongo': mongo_snapshot,
        'chains': chains,
        'triggers': triggers,
        'endpoints': endpoints,
        'counts': {'source_tables': len(tables), 'source_columns': sum(len(x['columns']) for x in tables.values()),
                   'source_fks': len(relations), 'source_enums': len(enums), 'source_excel_fields': len(source),
                   'master_tables': len(orm), 'master_business': len(set(orm) & BUSINESS),
                   'master_runtime': len(set(orm) - BUSINESS), 'master_fields': len(master),
                   'mongo_collections': len(mongo_snapshot['validators']),
                   'mongo_business_indexes': sum(map(len, mongo_snapshot['indexes'].values()))},
        'live': {'SOURCE': 'NOT_VERIFIED', 'MASTER': 'NOT_VERIFIED', 'MONGO': 'NOT_VERIFIED',
                 'reason': 'No explicitly configured independent test/live read-only target used; .env not loaded.'},
    }
    result = {name + '.csv': csv_text(rows) for name, rows in [
        ('source_erd_contract_matrix', source), ('source_enum_contract_matrix', enum_rows),
        ('source_relation_inventory', relations), ('master_mapping_matrix', master),
        ('mongo_contract_matrix', mongo), ('migration_trigger_inventory', migrations),
        ('trigger_inventory', triggers), ('endpoint_inventory', endpoints)]}
    result['contract_snapshot.json'] = json.dumps(snapshot, ensure_ascii=False, indent=2, sort_keys=True) + '\n'
    return result, snapshot


def main():
    sys.addaudithook(readonly_guard)
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--write', action='store_true', help='Generate derived audit artifacts under .test-tmp/audits/erd-contract only')
    mode.add_argument('--check', action='store_true', help='Compare artifacts with current inputs, no writes')
    args = parser.parse_args()
    artifacts, snapshot = collect()
    if args.write:
        OUTPUT.mkdir(parents=True, exist_ok=True)
        for name, body in artifacts.items():
            (OUTPUT / name).write_text(body, encoding='utf-8', newline='\n')
    if args.check:
        different = [name for name, body in artifacts.items() if not (OUTPUT / name).is_file()
                     or (OUTPUT / name).read_text(encoding='utf-8') != body]
        if different:
            raise SystemExit('Audit artifacts differ: ' + ', '.join(different))
    print(json.dumps({'mode': 'check' if args.check else 'write' if args.write else 'inspect',
                      'counts': snapshot['counts'], 'heads': {k: v['heads'] for k, v in snapshot['chains'].items()},
                      'artifact_count': len(artifacts)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
