"""Opt-in acceptance with a provisioned REAL local model and UUID test databases."""

import os
from dataclasses import replace
from hashlib import sha256
from math import isfinite

import pytest
from fastapi.testclient import TestClient
from pymongo import MongoClient
from sqlalchemy import text

from data_sync_etl.adapters.embeddings import LocalSentenceEmbedding
from data_sync_etl.adapters.mongo.content_store import MongoContentStore
from data_sync_etl.adapters.mongo.foundation_store import MongoBagStore, MongoContentReader
from data_sync_etl.adapters.mongo.schema import initialize
from data_sync_etl.adapters.ocr.engines import UnconfiguredOCREngine
from data_sync_etl.application.indexing import IndexingService
from data_sync_etl.config import Settings
from data_sync_etl.main import create_app
from tests.integration.source_postgres.test_file_import_flow import setup_import

pytestmark = [pytest.mark.postgres, pytest.mark.skipif(
    os.environ.get("ETL_TEST_REAL_EMBEDDING") != "1",
    reason="Opt-in local model acceptance: ETL_TEST_REAL_EMBEDDING=1 required",
)]
SAMPLES = (
    "Học sinh tìm hiểu định luật Newton và các ứng dụng trong môn Vật lý.",
    "Bài giảng Toán học giới thiệu phương trình bậc hai và phương pháp giải.",
)


def write_structured_docx(path, repetitions=8):
    from docx import Document
    document = Document()
    document.add_heading('Chương kiến thức tự nhiên', 1)
    document.add_heading('Bài học và vận dụng', 2)
    document.add_heading('Nội dung học tập', 3)
    for _ in range(repetitions):
        for sample in SAMPLES:
            document.add_paragraph(sample)
    document.save(path)


@pytest.fixture(scope="module")
def real_embedding():
    import torch

    config = Settings()
    torch.set_num_threads(2)
    model = LocalSentenceEmbedding(config.index_embedding_model_path,
                                   config.index_embedding_model_id,
                                   config.index_embedding_model_version,
                                   config.index_embedding_dimension)
    assert model.available(), "Provision the configured local model before this acceptance run"
    model.prepare()
    assert model.dimension == 384
    for sample in SAMPLES:
        vector = model.embed(sample)
        assert len(vector) == 384 and all(isfinite(v) for v in vector) and any(vector)
    return model


@pytest.fixture
def real_context(source, pg_container, mongo_test_database, tmp_path, real_embedding):
    c = pg_container
    client = MongoClient(mongo_test_database[0], serverSelectionTimeoutMS=3000, timeoutMS=10000)
    db = client[mongo_test_database[1]]
    assert db.name.startswith("etl_test_") and len(db.name) == 41
    assert not initialize(db, dimension=384)["ready"]
    assert initialize(db, apply=True, dimension=384)["ready"]
    c.settings.admin_enabled = True
    c.settings.index_embedding_dimension = 384
    c.settings.ocr_engine = "native"
    c.settings.chunk_size_words = 20
    c.settings.chunk_overlap_words = 4
    c.ocr.engine = UnconfiguredOCREngine()  # Real DOCX extraction, never fake OCR text/hierarchy.
    c.capabilities = replace(c.capabilities, ocr_engine="native", embedding_runtime=real_embedding)
    c.mongo_projector.documents = MongoContentStore(db, invalidate_content_bags=False)
    c.indexing = IndexingService(c.uow, MongoContentReader(db), real_embedding, c.settings, bags=MongoBagStore(db, 384))
    c.processing_finalizer.on_content_ready = c.indexing.enqueue_auto
    original = setup_import(source, c, tmp_path)
    original_bytes = original.read_bytes()
    content = original.with_suffix(".docx")
    write_structured_docx(content)
    with source.engine.begin() as connection:
        connection.execute(text("UPDATE resource_versions SET storage_object_key='native.docx',"
                                "source_hash=:hash,updated_at=now() WHERE id='source-local-1'"),
                           {"hash": sha256(content.read_bytes()).hexdigest()})
    try:
        yield c, db, source
        assert original.read_bytes() == original_bytes
        assert initialize(db, dimension=384)["ready"]
    finally:
        client.close()


def import_and_finish(c, expected_jobs=1):
    assert c.source_file_worker.once("real-model-file-test")
    for _ in range(50):
        if not c.worker.once("real-model-ocr-index-test"):
            break
    with c.uow() as repo:
        version = repo.find("resource_versions", order=("-version_number",), limit=1)[0]
        assert repo.count("ocr_jobs", {}) == expected_jobs
    assert c.source_file_import.list_files()["items"][0]["etl_status"] == "CONTENT_READY"
    return version


def test_manual_real_model_index_and_reindex(real_context, monkeypatch):
    c, db, _ = real_context
    version = import_and_finish(c)
    protected = {name: list(db[name].find({})) for name in ("users", "learning_resources", "ocr_pages")}
    index = c.indexing
    observed, inputs = [], []
    original_embed = index.embedding.embed
    def observe(value):
        with c.uow() as repo:
            observed.append(index._runs(repo, version["id"])[0]["status"])
        inputs.append(value)
        return original_embed(value)  # REAL runtime, no fallback.
    monkeypatch.setattr(index.embedding, "embed", observe)
    with TestClient(create_app(c)) as api:
        base = f"/api/v1/resource-versions/{version['id']}"
        assert api.get(base + "/chunks").json()["total"] > 10
        first = api.post(base + "/index")
        assert first.status_code == 202, first.text
        assert api.post(base + "/index").json()["action"] == "IN_PROGRESS"
        assert c.worker.once("real-model-index")
        status = api.get(base + "/index-status").json()
        assert status["index_status"] == "READY" and status["processing_status"] == "CONTENT_READY"
        assert status["bag_count"] == 1 and status["vector_ready"]
        docs = list(db.content_bags.find({}))
        assert len(docs) == 1 and set(observed) == {"RUNNING"}
        doc = docs[0]
        assert inputs == [doc["embedding_text"]]
        assert doc["embedding_text"] == " | ".join(doc["keywords"])
        assert doc["embedding_model"] == f"{index.embedding.model_id}@{index.embedding.version}"
        assert len(doc["embedding"]) == 384 and all(isfinite(v) for v in doc["embedding"]) and any(doc["embedding"])
        expected = original_embed(doc["embedding_text"])
        assert max(abs(a-b) for a,b in zip(expected, doc["embedding"], strict=True)) < 1e-6
        assert api.get(base + "/chunks?offset=10&limit=10").json()["items"][0]["chunk_sequence"] == 11
        second = api.post(base + "/reindex")
        assert second.status_code == 202 and second.json()["index_run_id"] != first.json()["index_run_id"]
        assert api.get(base + "/index-status").json()["published_run_id"] is None
        assert c.worker.once("real-model-reindex")
        assert db.content_bags.count_documents({}) == 1
        original_upsert = index.bags.upsert_bag
        def fail_storage(document, **kwargs):
            raise RuntimeError("Isolated test outage")
        monkeypatch.setattr(index.bags, "upsert_bag", fail_storage)
        assert api.post(base + "/reindex").status_code == 202
        assert c.worker.once("real-model-failure")
        failed = api.get(base + "/index-status").json()
        assert failed["index_status"] == "FAILED" and failed["processing_status"] == "CONTENT_READY"
        assert failed["published_run_id"] is None
        monkeypatch.setattr(index.bags, "upsert_bag", original_upsert)
        assert api.post(base + "/reindex").status_code == 202
        assert c.worker.once("real-model-recovery")
        assert index.status(version["id"])["index_status"] == "READY"
        with c.uow() as repo:
            assert repo.count("ocr_jobs", {}) == 1
        assert set(db.list_collection_names()) == {"users", "learning_resources", "ocr_pages", "mongo_content_units", "content_bags"}
        print("REAL_BAG_ACCEPTED keyword_embedding=PASS dimension=384 manual=PASS reindex=PASS")
    assert all(list(db[name].find({})) == docs for name, docs in protected.items())


def test_auto_real_model_source_to_index_no_backfill(real_context):
    c, db, source = real_context
    historical = import_and_finish(c)
    assert c.indexing.status(historical["id"])["index_status"] == "NOT_STARTED"
    c.settings.index_auto_enabled = True  # Test-only container; never changes the real .env.
    assert not c.worker.once("no-real-model-backfill")
    with c.uow() as repo:
        assert repo.count("etl_stage_runs", {"stage": "INDEX_BUILD"}) == 0
    updated = c.source_files.root / "books" / "updated.docx"
    write_structured_docx(updated, 9)
    with source.engine.begin() as connection:
        connection.execute(text(
            "INSERT INTO resource_versions (id,resource_id,parent_version_id,version_number,"
            "lifecycle_status,storage_bucket,storage_object_key,source_hash,uploaded_by,created_at,updated_at) "
            "SELECT 'source-local-2',resource_id,id,2,'APPROVED',storage_bucket,'updated.docx',"
            ":hash,uploaded_by,now(),now() FROM resource_versions WHERE id='source-local-1'"
        ), {"hash": sha256(updated.read_bytes()).hexdigest()})
    # Existing file discovery resumes after a Master Sync; exercise that real ordering.
    c.sync.run(full=False)
    c.source_file_import.discover_once()
    version = import_and_finish(c, expected_jobs=2)
    assert version["id"] != historical["id"]
    assert c.indexing.status(historical["id"])["index_status"] == "NOT_STARTED"
    state = c.indexing.status(version["id"])
    assert state["index_status"] == "READY" and state["vector_ready"] and state["chunk_count"] > 10
    with c.uow() as repo:
        assert repo.count("etl_stage_runs", {"stage": "INDEX_BUILD"}) == 1
        assert repo.count("ocr_jobs", {}) == 2
    c.sync.run(full=False)
    assert not c.source_file_worker.once("unchanged-source-real-model")
    assert not c.worker.once("unchanged-real-model-index")
    with c.uow() as repo:
        assert repo.count("etl_stage_runs", {"stage": "INDEX_BUILD"}) == 1
    assert db.content_bags.count_documents({}) == state["bag_count"] == 1
    c.settings.index_auto_enabled = False
    print(f"REAL_AUTO_ACCEPTED version={version['id']} chunks={state['chunk_count']} auto=PASS duplicate_prevention=PASS")


@pytest.mark.skipif(os.environ.get("ETL_TEST_REAL_UI") != "1", reason="Real browser acceptance is opt-in")
def test_react_ui_against_real_model_and_isolated_backend(real_context):
    import socket
    import subprocess
    from pathlib import Path
    from threading import Event, Thread
    from time import monotonic, sleep

    import uvicorn

    c, _, _ = real_context
    version = import_and_finish(c)
    c.indexing.enqueue(version["id"])
    assert c.worker.once("real-ui-initial-index")
    stop = Event()
    errors = []

    def worker_loop():
        while not stop.wait(0.1):
            try:
                c.worker.once("real-ui-existing-ocr-worker")
            except Exception as exc:
                errors.append(type(exc).__name__)
                return

    listener = socket.socket()
    listener.bind(("127.0.0.1", 0))
    listener.listen(128)
    port = listener.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(create_app(c), host="127.0.0.1", port=port,
                                          log_level="error", access_log=False))
    api_thread = Thread(target=lambda: server.run(sockets=[listener]), daemon=True)
    worker_thread = Thread(target=worker_loop, daemon=True)
    api_thread.start()
    worker_thread.start()
    try:
        deadline = monotonic() + 15
        while not server.started and monotonic() < deadline:
            sleep(0.05)
        assert server.started
        env = {**os.environ, "ETL_RUNTIME_API_URL": f"http://127.0.0.1:{port}",
               "ETL_RUNTIME_RESOURCE_ID": version["resource_id"], "ETL_RUNTIME_VERSION_ID": version["id"]}
        root = Path(__file__).parents[3]
        script = root / "frontend" / "tests" / "helpers" / "indexingRuntimeUi.mjs"
        assert script.is_file(), "Local real-browser acceptance harness must be present"
        result = subprocess.run(["node", str(script)], cwd=root, env=env, capture_output=True,
                                text=True, encoding="utf-8", timeout=90, check=False)
        assert result.returncode == 0, result.stdout + result.stderr
        assert not errors
        print(result.stdout.strip())
    finally:
        stop.set()
        server.should_exit = True
        worker_thread.join(timeout=10)
        api_thread.join(timeout=10)
        listener.close()
