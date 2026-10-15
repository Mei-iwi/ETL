from concurrent.futures import ThreadPoolExecutor
from hashlib import sha256

import pymupdf
import pytest
from conftest import migrate
from pymongo import MongoClient
from sqlalchemy import inspect, text

from data_sync_etl.adapters.mongo.content_store import MongoContentStore
from data_sync_etl.adapters.mongo.index_store import MongoIndexStore
from data_sync_etl.adapters.mongo.schema import initialize
from data_sync_etl.application.indexing import IndexingService
from data_sync_etl.application.sync import MasterSync

pytestmark = pytest.mark.postgres


@pytest.mark.parametrize('extension', ['.mp3', '.mp4'])
def test_source_media_excluded_from_import_queue(source, pg_container, tmp_path, extension):
    c = pg_container
    original = setup_import(source, c, tmp_path)
    with source.engine.begin() as connection:
        connection.execute(text('UPDATE resource_versions SET storage_object_key=:key, updated_at=now()'),
                           {'key': 'excluded' + extension})
    assert not c.source_file_worker.once('excluded-source')
    item = c.source_file_import.list_files()['items'][0]
    assert item['import_status'] == 'FAILED'
    assert item['last_error_code'] == 'UNSUPPORTED_SOURCE_FORMAT'
    assert item['retry_count'] == 0 and not item['retry_eligible']
    assert not c.source_file_worker.once('excluded-source')
    with c.uow() as repo:
        assert repo.count('resource_versions', {}) == repo.count('ocr_jobs', {}) == 0
    assert original.exists()


def test_import_queue_migration_up_down_up(pg_container):
    engine = pg_container.engine
    assert "source_file_imports" in inspect(engine).get_table_names()
    migrate(engine, "0010", downgrade=True)
    assert "source_file_imports" not in inspect(engine).get_table_names()
    migrate(engine)
    assert "source_file_imports" in inspect(engine).get_table_names()


def setup_import(source, c, tmp_path):
    root = tmp_path / "source"
    path = root / "books" / "native.pdf"
    path.parent.mkdir(parents=True)
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 72), "A genuine native text source PDF for integration testing.")
        doc.save(path)
    with source.engine.begin() as connection:
        connection.execute(text(f'GRANT SELECT ON "{source.schema}".resource_versions TO "{source.role}"'))
        user = connection.scalar(text("SELECT id FROM users ORDER BY id LIMIT 1"))
        connection.execute(text(
            "INSERT INTO resource_versions (id,resource_id,version_number,lifecycle_status,"
            "storage_bucket,storage_object_key,source_hash,uploaded_by,created_at,updated_at) "
            "VALUES ('source-local-1','resource-1',1,'APPROVED','books','native.pdf',:hash,:user,now(),now())"
        ), {"hash": sha256(path.read_bytes()).hexdigest(), "user": user})
    c.source = source.adapter
    c.source_file_import.source = source.adapter
    c.source_files.root = root
    c.sync = MasterSync(source.adapter, c.uow, batch_size=2)
    c.master_sync_worker.sync = c.sync
    c.sync.run(full=True)
    return path


def test_postgresql_concurrent_claim_and_mapping(source, pg_container, tmp_path):
    c = pg_container
    path = setup_import(source, c, tmp_path)
    original = path.read_bytes()
    c.source_file_import.discover_once()
    with ThreadPoolExecutor(max_workers=2) as pool:
        claims = list(pool.map(lambda _: c.source_file_import.claim(), range(2)))
    winners = [claim for claim in claims if claim is not None]
    assert len(winners) == 1
    c.source_file_import.execute(winners[0])
    c.sync.run(full=False)
    c.source_file_import.discover_once()
    with c.uow() as repo:
        assert repo.count("source_file_imports", {}) == repo.count("resource_versions", {}) == repo.count("ocr_jobs", {}) == 1
    assert original == path.read_bytes()
    assert c.source_file_import.list_files(failed_only=True)["items"] == []
    with c.uow() as repo:
        job = repo.find("ocr_jobs", limit=1)[0]
        repo.update("ocr_jobs", job["id"], {"status": "FAILED", "failed_pages": job["total_pages"]})
    assert len(c.source_file_import.list_files(failed_only=True)["items"]) == 1


def test_local_source_to_real_postgresql_and_mongo(source, pg_container, mongo_test_database, tmp_path):
    c = pg_container
    uri, database_name = mongo_test_database
    client = MongoClient(uri, serverSelectionTimeoutMS=3000)
    try:
        database = client[database_name]
        c.mongo_projector.documents = MongoContentStore(database)
        setup_import(source, c, tmp_path)
        # Exercise the existing Auto coordinator instead of asking a browser to schedule work.
        c.auto_sync.environment_supported = True
        c.auto_sync.source_health = source.adapter.health
        c.auto_sync.heartbeat_and_claim("master-test")
        c.auto_sync.configure(True, "resources")
        with c.uow() as repo:
            repo.update("sync_auto_config", "master-data", {"next_run_at": None})
        assert c.master_sync_worker.once("master-test")["action"] == "COMPLETED"
        assert c.source_file_worker.run("file-test", once=True)
        for _ in range(10):
            if not c.worker.once("ocr-test"):
                break
        files = c.source_file_import.list_files()
        assert files["items"][0]["etl_status"] == "CONTENT_READY"
        with c.uow() as repo:
            assert repo.count("content_units", {}) > 0
            assert repo.count("source_file_imports", {}) == repo.count("resource_versions", {}) == repo.count("ocr_jobs", {}) == 1
        assert database.ocr_pages.count_documents({}) == 1
        assert database.mongo_content_units.count_documents({}) > 0
        counts = (database.ocr_pages.count_documents({}), database.mongo_content_units.count_documents({}))
        c.sync.run(full=True)
        assert c.source_file_worker.run("file-test", once=True) is False
        assert counts == (database.ocr_pages.count_documents({}), database.mongo_content_units.count_documents({}))
    finally:
        client.close()


@pytest.mark.parametrize("auto_enabled", [False, True])
def test_source_file_to_index_with_existing_workers(source, pg_container, mongo_test_database, tmp_path, auto_enabled):
    class ModelDouble:
        model_id, version, dimension = "source-flow-test-double", "v1", 3

        def available(self):
            return True

        def prepare(self):
            pass

        def embed(self, text):
            return [0.1, 0.2, 0.3]

    c = pg_container
    client = MongoClient(mongo_test_database[0], serverSelectionTimeoutMS=3000, timeoutMS=10000)
    try:
        db = client[mongo_test_database[1]]
        initialize(db, apply=True)
        c.mongo_projector.documents = MongoContentStore(db, invalidate_content_bags=False)
        c.settings.index_auto_enabled = auto_enabled
        c.indexing = IndexingService(c.uow, MongoIndexStore(db), ModelDouble(), c.settings)
        c.processing_finalizer.on_content_ready = c.indexing.enqueue_auto
        original = setup_import(source, c, tmp_path)
        with c.uow() as repo:
            assert repo.count("etl_stage_runs", {"stage": "INDEX_BUILD"}) == 0
        assert c.source_file_worker.once("source-file-index-test")
        for _ in range(10):
            if not c.worker.once("ocr-index-test"):
                break
        with c.uow() as repo:
            version = repo.find("resource_versions", limit=1)[0]
            assert repo.count("etl_stage_runs", {"stage": "INDEX_BUILD"}) == int(auto_enabled)
            assert repo.count("ocr_jobs", {}) == 1
        assert c.indexing.status(version["id"])["index_status"] == ("READY" if auto_enabled else "NOT_STARTED")
        if not auto_enabled:
            c.indexing.enqueue(version["id"])
            assert c.worker.once("manual-index-test")
        assert db.indexed_chunks.count_documents({"resource_version_id": version["id"]}) > 0
        assert c.source_file_import.list_files()["items"][0]["etl_status"] == "CONTENT_READY"
        c.sync.run(full=False)
        assert not c.source_file_worker.once("same-source-test")
        assert original.exists()
    finally:
        client.close()


@pytest.mark.parametrize('extension', ['.docx', '.xlsx', '.txt', '.csv', '.pptx'])
def test_native_format_source_to_real_postgresql_mongo(source, pg_container, mongo_test_database, tmp_path, extension):
    c = pg_container
    original = setup_import(source, c, tmp_path)
    path = original.with_suffix(extension)
    if extension == '.docx':
        from docx import Document
        document = Document()
        document.add_heading('Bài học phân số', 1)
        document.add_paragraph('Phép cộng và phép trừ')
        document.save(path)
    elif extension == '.xlsx':
        from openpyxl import Workbook
        document = Workbook()
        document.active.title = 'Toán'
        document.active.append(['Phân số', 2])
        document.create_sheet('Văn').append(['Tiếng Việt'])
        document.save(path)
    elif extension == '.pptx':
        from pptx import Presentation
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[1]).shapes.title.text = 'Bài học phân số'
        document.save(path)
    else:
        path.write_text('Bài học phân số\nPhép cộng và phép trừ' if extension == '.txt'
                        else 'Tên,Nội dung\nPhân số,"Phép cộng, phép trừ"', encoding='utf-8')
    with source.engine.begin() as connection:
        connection.execute(text("UPDATE resource_versions SET storage_object_key=:key, source_hash=:hash, updated_at=now()"),
                           {'key': path.name, 'hash': sha256(path.read_bytes()).hexdigest()})
    uri, name = mongo_test_database
    with MongoClient(uri, serverSelectionTimeoutMS=3000) as client:
        database = client[name]
        c.mongo_projector.documents = MongoContentStore(database)
        assert c.source_file_worker.once('source-native')
        for _ in range(10):
            if not c.worker.once('native-test'):
                break
        file = c.source_file_import.list_files()['items'][0]
        assert file['etl_status'] == 'CONTENT_READY'
        with c.uow() as repo:
            units = repo.find('content_units', {'resource_version_id': file['target_version_id']})
            assert units and all(unit['page_from'] is None and unit['title'] for unit in units)
            assert repo.count('resource_versions', {}) == repo.count('ocr_jobs', {}) == 1
        projected = list(database.mongo_content_units.find({'resource_version_id': file['target_version_id']}))
        assert len(projected) == len(units) and any('phân số' in unit['text'].lower() for unit in projected)
        assert all(unit['page_from'] is None and unit['title'] for unit in projected)
        before = {unit['_id'] for unit in projected}
        c.sync.run(full=True)
        assert not c.source_file_worker.once('source-native')
        assert before == {unit['_id'] for unit in database.mongo_content_units.find({})}
