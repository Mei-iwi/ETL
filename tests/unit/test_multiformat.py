import io
from datetime import date
from hashlib import sha256
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from data_sync_etl.adapters.extraction import NativeDocumentExtractor
from data_sync_etl.domain.core import DomainError, now
from data_sync_etl.domain.format_validation import FormatValidatorRegistry, InvalidFileContent
from data_sync_etl.main import create_app
from data_sync_etl.ports.processing import OCRResult
from tests.indexing_support import setup_index
from tests.local_admin_client import TestClient


def sample(extension):
    output = io.BytesIO()
    if extension == '.docx':
        from docx import Document
        document = Document()
        document.add_heading('Bài học phân số', 1)
        document.add_paragraph('Đoạn một: phép cộng.')
        document.add_paragraph('Đoạn hai: phép trừ.')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = 'Tên', 'Giá trị'
        table.cell(1, 0).text, table.cell(1, 1).text = 'Phân số', '1/2'
        document.save(output)
    elif extension == '.xlsx':
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Toán'
        sheet.append(['Học liệu', 'Số lượng', 'Ngày', 'Duyệt'])
        sheet.append(['Phân số', 2, date(2026, 10, 14), True])
        sheet.append(['Công thức', '=1+2'])
        workbook.create_sheet('Văn').append(['Tiếng Việt', 'Dấu tiếng Việt'])
        workbook.save(output)
    elif extension == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        first = document.slides.add_slide(document.slide_layouts[1])
        first.shapes.title.text = 'Bài học phân số'
        first.placeholders[1].text = 'Phép cộng và phép trừ'
        second = document.slides.add_slide(document.slide_layouts[5])
        second.shapes.title.text = 'Ví dụ'
        table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(2)).table
        table.cell(0, 0).text, table.cell(0, 1).text = 'Tên', 'Giá trị'
        table.cell(1, 0).text, table.cell(1, 1).text = 'Phân số', '1/2'
        document.save(output)
    elif extension in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (1000, 200), 'white')
        try:
            font = ImageFont.truetype('arial.ttf', 40)
        except OSError:
            font = ImageFont.load_default(size=40)
        ImageDraw.Draw(image).text((25, 60), 'Bài học phân số', font=font, fill='black')
        image.save(output, format='PNG' if extension == '.png' else 'JPEG')
    elif extension == '.txt':
        return '\ufeffBài học phân số\nPhép cộng và phép trừ\n'.encode()
    elif extension == '.csv':
        return '\ufeffTên,Nội dung\r\nPhân số,"Phép cộng, phép trừ\nDòng tiếp theo"\r\n'.encode()
    elif extension == '.pdf':
        import pymupdf
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), 'Native PDF regression sample')
            return document.tobytes()
    else:
        raise ValueError(extension)
    return output.getvalue()


class ImageEngineDouble:
    """Logic test only: not evidence of real Tesseract recognition."""
    supports_images = True

    def check_runtime(self):
        pass

    def recognize(self, path):
        assert path.exists()
        return OCRResult('Bài học phân số', engine_name='image_test_double')


FORMATS = ['.pdf', '.docx', '.xlsx', '.txt', '.csv', '.png', '.jpg', '.jpeg', '.pptx']


def xlsx_with_default_content_type(*, override=None, default_type=None, duplicate_default=False,
                                   missing_workbook=False):
    """Valid XLSX variant exported by applications using OPC Default mappings."""
    validator = FormatValidatorRegistry().find_validator('sample.xlsx')
    namespace = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    output = io.BytesIO()
    with ZipFile(io.BytesIO(sample('.xlsx'))) as original, ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for item in original.infolist():
            if missing_workbook and item.filename == validator.main_part:
                continue
            payload = original.read(item)
            if item.filename == '[Content_Types].xml':
                root = ElementTree.fromstring(payload)
                for node in list(root):
                    if node.get('PartName') == '/' + validator.main_part:
                        root.remove(node)
                    elif node.get('Extension') == 'xml':
                        node.set('ContentType', default_type or validator.content_type)
                if override is not None:
                    ElementTree.SubElement(root, namespace + 'Override',
                                           PartName='/' + validator.main_part, ContentType=override)
                if duplicate_default:
                    ElementTree.SubElement(root, namespace + 'Default', Extension='xml',
                                           ContentType=validator.content_type)
                payload = ElementTree.tostring(root, encoding='utf-8', xml_declaration=True)
            archive.writestr(item, payload)
    return output.getvalue()


def test_xlsx_default_content_type_validates_and_extracts(tmp_path):
    path = tmp_path / 'default.xlsx'
    path.write_bytes(xlsx_with_default_content_type())
    FormatValidatorRegistry().find_validator(path.name).validate_path(path)
    result = NativeDocumentExtractor().extract(path)
    assert 'Phân số' in result.text and result.metadata['source_type'] == 'xlsx'


@pytest.mark.parametrize('options', [
    {'default_type': 'application/xml'},
    {'default_type': 'application/vnd.ms-excel.sheet.macroEnabled.main+xml'},
    {'override': 'application/vnd.ms-excel.sheet.macroEnabled.main+xml'},
    {'override': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'},
    {'duplicate_default': True},
    {'missing_workbook': True},
])
def test_xlsx_default_content_type_does_not_bypass_mismatch_checks(tmp_path, options):
    path = tmp_path / 'invalid.xlsx'
    path.write_bytes(xlsx_with_default_content_type(**options))
    with pytest.raises(InvalidFileContent, match='OPENXML_TYPE_MISMATCH'):
        FormatValidatorRegistry().find_validator(path.name).validate_path(path)


def test_xlsx_override_takes_precedence_over_default(tmp_path):
    path = tmp_path / 'override.xlsx'
    validator = FormatValidatorRegistry().find_validator(path.name)
    path.write_bytes(xlsx_with_default_content_type(override=validator.content_type, default_type='application/xml'))
    validator.validate_path(path)


@pytest.mark.parametrize('extension', FORMATS)
def test_real_upload_worker_content_units_with_memory_projection(container, extension):
    setup_index(container)
    container.settings.admin_enabled = True
    if extension in {'.png', '.jpg', '.jpeg'}:
        container.ocr.engine = ImageEngineDouble()
    validator = FormatValidatorRegistry().find_validator('sample' + extension)
    with TestClient(create_app(container), client=('127.0.0.1', 50000)) as client:
        response = client.post('/api/v1/resources', data={
            'title': 'Format test', 'resource_type_code': 'HANDOUT',
            'provider_id': 'provider-test', 'provider_name': 'Test provider',
        }, files={'file': ('sample' + extension, sample(extension), sorted(validator.mime_types)[0])})
        assert response.status_code == 201, response.text
        version_id = response.json()['initial_version_id']
        started = client.post(f'/api/v1/resource-versions/{version_id}/process')
        assert started.status_code == 202, started.text
        job_id = started.json()['job_id']
        for _ in range(20):
            if not container.worker.once('format-test'):
                break
        status = client.get(f'/api/v1/processing-jobs/{job_id}').json()
        assert status['processing_status'] == 'CONTENT_READY', status
        units = client.get(f'/api/v1/resource-versions/{version_id}/content-units').json()['items']
        assert units
        if extension != '.pdf':
            assert any('phân số' in unit['text'].lower() for unit in units)
        if extension != '.pdf':
            assert all(unit['page_from'] is None and unit['page_to'] is None and unit['title'] for unit in units)
            assert all(unit['unit_type'] != 'PAGE_TEXT' for unit in units)
        assert len(container.documents.units) == len(units)
        again = client.post(f'/api/v1/resource-versions/{version_id}/process').json()
        assert again['job_id'] == job_id and again['processing_status'] == 'CONTENT_READY'
        container.postprocess.run(version_id)
        assert len(client.get(f'/api/v1/resource-versions/{version_id}/content-units').json()['items']) == len(units)


@pytest.mark.parametrize('extension', ['.docx', '.xlsx', '.txt', '.csv', '.pptx'])
def test_extract_native_provenance_and_unicode(tmp_path, extension):
    path = tmp_path / ('sample' + extension)
    path.write_bytes(sample(extension))
    result = NativeDocumentExtractor().extract(path)
    assert 'phân số' in result.text.lower()
    assert result.metadata['source_type'] == extension[1:]
    units = result.metadata['logical_units']
    assert all(result.text[unit['start']:unit['end']] and unit['locator'] for unit in units)
    if extension == '.docx':
        assert result.text.index('Đoạn một') < result.text.index('Đoạn hai') < result.text.index('Giá trị')
    if extension == '.xlsx':
        assert {unit['locator']['sheet'] for unit in units} == {'Toán', 'Văn'}
        assert 'A2: Phân số' in result.text and '2026-10-14' in result.text
        assert 'Công thức chưa có giá trị cache' in result.text and '=1+2' not in result.text
    if extension == '.csv':
        assert 'Phép cộng, phép trừ\nDòng tiếp theo' in result.text
        assert units[-1]['locator']['line_to'] == 3
    if extension == '.pptx':
        assert {unit['locator']['slide'] for unit in units} == {1, 2}


@pytest.mark.parametrize('extension', ['.docx', '.xlsx', '.png', '.csv'])
def test_source_import_preserves_original_bytes_and_projection(container, tmp_path, monkeypatch, extension):
    container.sync.run(full=True)
    container.settings.source_reference_enabled = True
    container.settings.source_temporary_files_enabled = True
    container.settings.source_temp_root = tmp_path / 'source-temp'
    with container.uow() as repo:
        repo.update('master_learning_resources', 'resource-1', {'source_origin': 'POSTGRESQL'})
    root = tmp_path / 'source'
    path = root / 'books' / ('sample' + extension)
    path.parent.mkdir(parents=True)
    payload = sample(extension)
    path.write_bytes(payload)
    container.source_files.root = root
    if extension == '.png':
        container.ocr.engine = ImageEngineDouble()
    rows = [{'source_version_id': 'source-format', 'resource_id': 'resource-1',
             'resource_title': 'Format import', 'source_version_number': 1, 'parent_version_id': None,
             'source_lifecycle_status': 'APPROVED', 'updated_at': now(),
             'source_hash': sha256(payload).hexdigest(), 'storage_bucket': 'books',
             'storage_object_key': path.name}]
    monkeypatch.setattr(container.source, 'version_page', lambda **_: {'items': rows, 'next_cursor': None}, raising=False)
    monkeypatch.setattr(container.source, 'get_version', lambda id: next(
        (dict(r) for r in rows if r['source_version_id'] == id), None), raising=False)
    assert container.source_file_worker.once('source-format')
    for _ in range(20):
        if not container.worker.once('format-test'):
            break
    listed = container.source_file_import.list_files()['items'][0]
    assert listed['etl_status'] == 'CONTENT_READY'
    with container.uow() as repo:
        version = repo.get('resource_versions', listed['target_version_id'])
        assert repo.count('resource_versions', {}) == 1
    with container.source_resolver.open_verified_stream(container.source_resolver.resolve(version['id'])) as stream:
        assert stream.read() == payload
    assert not list((container.storage.root / 'resources').rglob('*'))
    container.source_file_import.discover_once()
    assert not container.source_file_worker.once('source-format')


@pytest.mark.parametrize('extension,mime,data', [
    ('.txt', 'text/plain', b'\xff\xfe\x00'),
    ('.csv', 'text/csv', b'one\x00,two'),
    ('.png', 'image/png', b'\x89PNG\r\n\x1a\nnot-an-image'),
    ('.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', b'PK\x03\x04broken'),
])
def test_invalid_content_rejected(tmp_path, extension, mime, data):
    validator = FormatValidatorRegistry().find_validator('sample' + extension)
    path = tmp_path / ('sample' + extension)
    path.write_bytes(data)
    with pytest.raises(DomainError):
        validator.validate(extension, mime, data[:16], len(data))
        validator.validate_path(path)


def test_zip_type_bomb_and_traversal_rejected(tmp_path):
    validator = FormatValidatorRegistry().find_validator('sample.docx')
    path = tmp_path / 'sample.docx'
    for name, data in [('xl/workbook.xml', b'<root/>'), ('../evil.xml', b'<root/>'),
                       ('word/document.xml', b'A' * (2 * 1024 * 1024))]:
        with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
            archive.writestr(name, data)
        with pytest.raises(InvalidFileContent):
            validator.validate_path(path)
    path.write_bytes(sample('.xlsx'))
    with pytest.raises(InvalidFileContent, match='TYPE_MISMATCH'):
        validator.validate_path(path)


def test_text_empty_limit_csv_semicolon_and_invalid_csv(tmp_path, monkeypatch):
    path = tmp_path / 'sample.csv'
    path.write_text('Tên;Nội dung\nPhân số;"Phép cộng; phép trừ"\n', encoding='utf-8')
    assert 'Phép cộng; phép trừ' in NativeDocumentExtractor().extract(path).text
    path.write_text('Tên,Nội dung\n"unterminated', encoding='utf-8')
    with pytest.raises(DomainError, match='INVALID_CSV'):
        NativeDocumentExtractor().extract(path)
    path = tmp_path / 'empty.txt'
    path.write_text(' \n\n', encoding='utf-8')
    with pytest.raises(DomainError, match='EMPTY_EXTRACTED_CONTENT'):
        NativeDocumentExtractor().extract(path)
    monkeypatch.setattr('data_sync_etl.adapters.extraction.MAX_TEXT_CHARS', 5)
    path.write_text('Bài học phân số', encoding='utf-8')
    with pytest.raises(DomainError, match='LIMIT_EXCEEDED'):
        NativeDocumentExtractor().extract(path)


@pytest.mark.parametrize('extension', FORMATS)
def test_validator_mime_empty_and_size_limits(extension):
    validator = FormatValidatorRegistry().find_validator('sample' + extension)
    payload = sample(extension)
    with pytest.raises(DomainError):
        validator.validate(validator.extension, 'application/octet-stream', payload[:16], len(payload))
    for size in (0, validator.max_size_bytes + 1):
        with pytest.raises(DomainError):
            validator.validate(validator.extension, sorted(validator.mime_types)[0], payload[:16], size)


def test_image_pixel_limit_and_utf8_without_bom(tmp_path, monkeypatch):
    path = tmp_path / 'sample.png'
    path.write_bytes(sample('.png'))
    monkeypatch.setattr('data_sync_etl.domain.format_validation.MAX_IMAGE_PIXELS', 1)
    with pytest.raises(InvalidFileContent, match='INVALID_IMAGE'):
        FormatValidatorRegistry().find_validator(path.name).validate_path(path)
    text = tmp_path / 'sample.txt'
    text.write_text('Tiếng Việt\nCó dấu', encoding='utf-8')
    assert 'Tiếng Việt' in NativeDocumentExtractor().extract(text).text


@pytest.mark.parametrize('extension', ['.png', '.jpg', '.jpeg'])
def test_real_tesseract_image_runtime_if_installed(tmp_path, extension):
    from data_sync_etl.adapters.extraction import recognize_image
    from data_sync_etl.adapters.ocr.engines import TesseractOCREngine
    engine = TesseractOCREngine()
    try:
        engine.check_runtime()
    except DomainError:
        pytest.skip('Real Tesseract with vie+eng is not installed/configured; OCR quality NOT VERIFIED')
    path = tmp_path / ('vietnamese' + extension)
    path.write_bytes(sample(extension))
    result, _ = recognize_image(path, tmp_path / 'oriented.png', engine)
    assert result.engine_name == 'tesseract' and result.text.strip()
    assert 'phân' in result.text.lower() or 'bài' in result.text.lower()


def test_blank_csv_and_formula_only_workbook_have_no_fake_content(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / 'empty.csv'
    path.write_text(',\n,\n', encoding='utf-8')
    with pytest.raises(DomainError, match='EMPTY_EXTRACTED_CONTENT'):
        NativeDocumentExtractor().extract(path)
    path = tmp_path / 'formula.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = '=1+2'
    workbook.save(path)
    with pytest.raises(DomainError, match='EMPTY_EXTRACTED_CONTENT'):
        NativeDocumentExtractor().extract(path)


def test_openxml_rejects_dtd_and_macro(tmp_path):
    path = tmp_path / 'sample.docx'
    validator = FormatValidatorRegistry().find_validator(path.name)
    for malicious, payload in [('word/evil.xml', b'<!DOCTYPE x [<!ENTITY t "text">]><x>&t;</x>'),
                               ('word/vbaProject.bin', b'ignored-macro')]:
        with ZipFile(io.BytesIO(sample('.docx'))) as original, ZipFile(path, 'w') as archive:
            for item in original.infolist():
                archive.writestr(item, original.read(item))
            archive.writestr(malicious, payload)
        with pytest.raises(InvalidFileContent):
            validator.validate_path(path)


def test_xlsx_dimensions_are_bounded(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / 'wide.xlsx'
    workbook = Workbook()
    workbook.active['XFD1'] = 'Too wide'
    workbook.save(path)
    with pytest.raises(DomainError, match='EXTRACTION_LIMIT_EXCEEDED'):
        NativeDocumentExtractor().extract(path)


def test_xlsx_cached_formula_is_read_not_evaluated(tmp_path):
    source = io.BytesIO(sample('.xlsx'))
    output = io.BytesIO()
    with ZipFile(source) as original, ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for item in original.infolist():
            data = original.read(item)
            if item.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<f>1+2</f><v></v>', b'<f>1+2</f><v>3</v>')
            archive.writestr(item, data)
    path = tmp_path / 'cache.xlsx'
    path.write_bytes(output.getvalue())
    result = NativeDocumentExtractor().extract(path)
    assert 'B3: 3' in result.text and '=1+2' not in result.text


def test_empty_native_document_fails_without_units(container):
    container.settings.admin_enabled = True
    with TestClient(create_app(container), client=('127.0.0.1', 50000)) as client:
        created = client.post('/api/v1/resources', data={
            'title': 'Empty text', 'resource_type_code': 'HANDOUT',
            'provider_id': 'provider-test', 'provider_name': 'Test provider',
        }, files={'file': ('sample.txt', b' \n', 'text/plain')}).json()
        started = client.post(f"/api/v1/resource-versions/{created['initial_version_id']}/process").json()
        container.worker.once('empty-text')
        status = client.get(f"/api/v1/processing-jobs/{started['job_id']}").json()
        assert status['processing_status'] == 'OCR_FAILED'
    assert not container.documents.units


def test_native_processing_retry_keeps_version_job_and_units(container, monkeypatch):
    container.settings.admin_enabled = True
    mime = sorted(FormatValidatorRegistry().find_validator('sample.docx').mime_types)[0]
    extractor = container.ocr.native_documents.extract
    with TestClient(create_app(container), client=('127.0.0.1', 50000)) as client:
        created = client.post('/api/v1/resources', data={
            'title': 'Retry document', 'resource_type_code': 'HANDOUT',
            'provider_id': 'provider-test', 'provider_name': 'Test provider',
        }, files={'file': ('sample.docx', sample('.docx'), mime)}).json()
        version_id = created['initial_version_id']
        started = client.post(f'/api/v1/resource-versions/{version_id}/process').json()
        job_id = started['job_id']
        def unavailable(path):
            raise RuntimeError('Temporary extractor failure')
        monkeypatch.setattr(container.ocr.native_documents, 'extract', unavailable)
        for _ in range(3):
            container.worker.once('native-retry')
        assert client.get(f'/api/v1/processing-jobs/{job_id}').json()['processing_status'] == 'OCR_FAILED'
        monkeypatch.setattr(container.ocr.native_documents, 'extract', extractor)
        queued = client.post(f'/api/v1/processing-jobs/{job_id}/retry')
        assert queued.status_code == 202 and queued.json()['job_id'] == job_id
        for _ in range(10):
            if not container.worker.once('native-retry'):
                break
        assert client.get(f'/api/v1/processing-jobs/{job_id}').json()['processing_status'] == 'CONTENT_READY'
        with container.uow() as repo:
            assert repo.count('resource_versions', {}) == repo.count('ocr_jobs', {}) == 1
            assert repo.count('content_units', {}) == len(container.documents.units)


@pytest.mark.parametrize('extension', ['.docx', '.xlsx', '.txt', '.csv', '.pptx'])
def test_update_version_accepts_native_formats_without_duplicate_versions(container, extension):
    container.settings.admin_enabled = True
    container.sync.run(full=True)
    with container.uow() as repo:
        repo.update("master_learning_resources", "resource-1", {"source_origin": "LOCAL"})
    mime = sorted(FormatValidatorRegistry().find_validator('sample' + extension).mime_types)[0]
    payload = sample(extension)
    with TestClient(create_app(container), client=('127.0.0.1', 50000)) as client:
        first = client.post('/api/v1/resource/resource-1/versions',
                            files={'file': ('sample' + extension, payload, mime)})
        assert first.status_code == 201, first.text
        again = client.post('/api/v1/resource/resource-1/versions',
                            files={'file': ('sample' + extension, payload, mime)})
        assert again.status_code == 200 and again.json()['action'] == 'NO_CHANGE'
        assert again.json()['resource_version_id'] == first.json()['resource_version_id']


@pytest.mark.parametrize('extension,payload,expected', [
    ('.docx', b'PK\x03\x04broken', 'INVALID_OPENXML_CONTAINER'),
])
def test_source_invalid_formats_fail_without_target_version(container, tmp_path, monkeypatch, extension, payload, expected):
    container.sync.run(full=True)
    container.settings.source_reference_enabled = True
    container.settings.source_temporary_files_enabled = True
    container.settings.source_temp_root = tmp_path / 'source-temp'
    with container.uow() as repo:
        repo.update('master_learning_resources', 'resource-1', {'source_origin': 'POSTGRESQL'})
    root = tmp_path / 'source'
    path = root / 'books' / ('sample' + extension)
    path.parent.mkdir(parents=True)
    path.write_bytes(payload)
    container.source_files.root = root
    rows = [{'source_version_id': 'bad-format', 'resource_id': 'resource-1',
             'resource_title': 'Bad import', 'source_version_number': 1, 'parent_version_id': None,
             'source_lifecycle_status': 'APPROVED', 'updated_at': now(),
             'source_hash': sha256(payload).hexdigest(), 'storage_bucket': 'books',
             'storage_object_key': path.name}]
    monkeypatch.setattr(container.source, 'version_page', lambda **_: {'items': rows, 'next_cursor': None}, raising=False)
    monkeypatch.setattr(container.source, 'get_version', lambda id: next(
        (dict(r) for r in rows if r['source_version_id'] == id), None), raising=False)
    assert container.source_file_worker.once('bad-format')
    listed = container.source_file_import.list_files()['items'][0]
    assert listed['import_status'] == 'FAILED' and listed['last_error_code'] == expected
    with container.uow() as repo:
        assert repo.count('resource_versions', {}) == repo.count('ocr_jobs', {}) == 0


@pytest.mark.parametrize('extension,payload,mime', [('.png', sample('.png'), 'image/png')])
def test_process_rejects_unsupported_or_missing_runtime_without_job(container, extension, payload, mime):
    container.settings.admin_enabled = True
    with TestClient(create_app(container), client=('127.0.0.1', 50000)) as client:
        response = client.post('/api/v1/resources', data={
            'title': 'Unsupported test', 'resource_type_code': 'HANDOUT',
            'provider_id': 'provider-test', 'provider_name': 'Test provider',
        }, files={'file': ('sample' + extension, payload, mime)})
        assert response.status_code == 201, response.text
        rejected = client.post(f"/api/v1/resource-versions/{response.json()['initial_version_id']}/process")
        assert rejected.status_code == 400
        assert rejected.json()['detail']['code'] == ('BLOCKED_OCR_RUNTIME' if extension == '.png' else 'UNSUPPORTED_PROCESSING_FORMAT')
    with container.uow() as repo:
        assert repo.count('ocr_jobs', {}) == 0
