"""Real DOCX Heading/OOXML -> existing extractor/postprocess, not seeded PG trees."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from data_sync_etl.adapters.extraction import NativeDocumentExtractor
from data_sync_etl.application.postprocess import Postprocessor
from tests.indexing_support import setup_index


def finish_docx(c, levels=(1, 2, 3), auto=False):
    c.sync.run(full=True)
    path = c.storage.get_local_path('input', 'headings.docx')
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for level in levels:
        document.add_heading('Tiêu đề cấp ' + str(level), level)
    document.add_paragraph('Học sinh nghiên cứu định luật vật lý và toán học.')
    document.add_paragraph('Phương trình toán học giúp giải bài tập thực nghiệm.')
    document.save(path)
    c.settings.index_auto_enabled = auto
    version = c.ingestion.run('resource-1', 'input', 'headings.docx')['version']
    c.ocr.create_job(version['id'])
    while c.worker.once('docx-test'):
        pass
    return version


@pytest.mark.parametrize('levels,root', [((1, 2, 3), 'CHAPTER'), ((2, 3), 'LESSON'), ((3,), 'SECTION')])
def test_docx_real_hierarchy_auto_index_idempotent(container, levels, root):
    index = setup_index(container)
    version = finish_docx(container, levels, True)
    status = index.status(version['id'])
    assert status['processing_status'] == 'CONTENT_READY'
    assert status['index_status'] == 'READY', status
    assert status['bag_count'] == 1 and status['bags'][0]['unit_type'] == root
    with container.uow() as repo:
        units = repo.find('content_units', {'resource_version_id': version['id']}, order=('sequence_no',))
        ids = [u['id'] for u in units]
        assert units[0]['parent_unit_id'] is None
        assert all(u['parent_unit_id'] is not None for u in units[1:])
        assert all(u['resource_version_id'] == version['id'] for u in units)
        assert len(index._runs(repo, version['id'])) == 1
    # Same results/outline produce the same PG ContentUnit identities and flags.
    container.postprocess.run(version['id'])
    with container.uow() as repo:
        rerun = repo.find('content_units', {'resource_version_id': version['id']}, order=('sequence_no',))
        assert [u['id'] for u in rerun] == ids
        assert [u['parent_unit_id'] for u in rerun] == [u['parent_unit_id'] for u in units]
    previews = index.chunks(version['id'])['items']
    assert len(previews) == len(levels) + 2
    assert len({p['text'] for p in previews}) == len(previews)


def test_unsupported_docx_outline_stays_flat(container):
    index = setup_index(container)
    version = finish_docx(container, (1, 4), True)
    status = index.status(version['id'])
    assert status['processing_status'] == 'CONTENT_READY'
    assert status['index_status'] == 'BLOCKED_STRUCTURE'
    assert index.chunks(version['id'])['total'] == 4
    assert not index.bags.rows


def test_extractor_uses_outline_not_title_text(tmp_path):
    document = Document()
    document.add_paragraph('CHƯƠNG 1 giả bằng chữ nhưng không có Heading style')
    paragraph = document.add_paragraph('Đề mục thật')
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '2')
    paragraph._p.get_or_add_pPr().append(outline)
    path = tmp_path / 'outline.docx'
    document.save(path)
    result = NativeDocumentExtractor().extract(path)
    assert 'heading_level' not in result.metadata['logical_units'][0]['locator']
    assert result.metadata['logical_units'][1]['locator']['heading_level'] == 3
    assert result.metadata['structure_version'] == 'docx-outline-v1'


def test_malformed_outline_metadata_never_becomes_root():
    result = dict(raw_text='Nội dung', result_metadata=dict(source_type='docx', structure_version='docx-outline-v1',
        logical_units=[dict(start=0, end=8, title='T', locator={'heading_level': True})]))
    units = list(Postprocessor._text_units(result, 1))
    assert units[0][3] == 'DOCX_TEXT' and units[0][4] is None


def test_existing_worker_gives_indexing_a_fair_turn():
    from types import SimpleNamespace

    from data_sync_etl.workers.ocr import OCRWorker

    reconciled = []
    pipeline = SimpleNamespace(heartbeat_timeout=300,
        recover_stale=lambda: {'recovered': 0}, claim=lambda _: {'id': 't', 'job_id': 'j'},
        recognize=lambda _: (None, None), succeed=lambda *args: None)
    worker = OCRWorker(pipeline, reconcile=lambda: reconciled.append(True) or True)
    for _ in range(5):
        assert worker.once('test')
    assert not reconciled
    assert worker.once('test') and len(reconciled) == 1


def test_empty_heading_does_not_make_bag_ready(container):
    index = setup_index(container)
    container.sync.run(full=True)
    path = container.storage.get_local_path('input', 'heading-only.docx')
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading('Chỉ có tiêu đề', 1)
    document.save(path)
    version = container.ingestion.run('resource-1', 'input', path.name)['version']
    container.ocr.create_job(version['id'])
    while container.worker.once('empty-heading'):
        pass
    assert index.status(version['id'])['index_status'] == 'BLOCKED_STRUCTURE'
    assert not index.bags.rows


def test_outline_nodes_only_are_not_content_leaves():
    from data_sync_etl.application.content_bags import bag_structure
    from data_sync_etl.domain.core import DomainError
    rows = [dict(id='chapter', resource_version_id='v', parent_unit_id=None, unit_type='CHAPTER'),
            dict(id='section', resource_version_id='v', parent_unit_id='chapter', unit_type='SECTION')]
    with pytest.raises(DomainError, match='descendant content leaves'):
        bag_structure(rows)
